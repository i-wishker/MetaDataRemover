import os
from tkinter import Tk
from tkinter.filedialog import askopenfilename
from mutagen import File as MutagenFile
from pyexiv2 import Image
from docx import Document  # Import for handling .docx files
from pptx import Presentation  # Import for handling .pptx files
from openpyxl import load_workbook  # Import for handling .xlsx files

def remove_metadata(file_path):
    try:
        # Handle image files
        if file_path.lower().endswith(('.jpg', '.jpeg', '.png', '.tiff', '.bmp')):
            with Image(file_path) as img:
                img.clear_metadata()
                img.save()
            print(f"Metadata cleared for image: {file_path}")
        
        # Handle audio files
        elif file_path.lower().endswith(('.mp3', '.flac', '.wav', '.ogg', '.m4a')):
            audio = MutagenFile(file_path, easy=True)
            if audio:
                audio.delete()
                audio.save()
            print(f"Metadata cleared for audio: {file_path}")
        
        # Handle .docx files
        elif file_path.lower().endswith('.docx'):
            doc = Document(file_path)
            if doc.core_properties:
                doc.core_properties.author = None
                doc.core_properties.title = None
                doc.core_properties.subject = None
                doc.core_properties.comments = None
                doc.core_properties.keywords = None
                doc.core_properties.last_modified_by = None
                doc.core_properties.revision = 1  # Set to 1 instead of 0
            doc.save(file_path)
            print(f"Metadata cleared for Word document: {file_path}")
        
        # Handle .pptx files
        elif file_path.lower().endswith('.pptx'):
            ppt = Presentation(file_path)
            if ppt.core_properties:
                ppt.core_properties.author = None
                ppt.core_properties.title = None
                ppt.core_properties.subject = None
                ppt.core_properties.comments = None
                ppt.core_properties.keywords = None
                ppt.core_properties.last_modified_by = None
                ppt.core_properties.revision = 1  # Set to 1 instead of 0
            ppt.save(file_path)
            print(f"Metadata cleared for PowerPoint presentation: {file_path}")
        
        # Handle .xlsx files
        elif file_path.lower().endswith('.xlsx'):
            wb = load_workbook(file_path)
            if wb.properties:
                wb.properties.creator = None
                wb.properties.title = None
                wb.properties.subject = None
                wb.properties.description = None
                wb.properties.keywords = None
                wb.properties.lastModifiedBy = None
                wb.properties.revision = 1  # Set to 1 instead of 0
            wb.save(file_path)
            print(f"Metadata cleared for Excel spreadsheet: {file_path}")
        
        # Add more file types as needed
        else:
            print(f"Unsupported file type: {file_path}")
    except Exception as e:
        print(f"Error clearing metadata for {file_path}: {e}")

if __name__ == "__main__":
    # Create a Tkinter root window and hide it
    root = Tk()
    root.withdraw()
    root.title("Select File")

    # Open a file selection dialog
    file_path = askopenfilename(title="Select File to Clear Metadata")
    if file_path:
        remove_metadata(file_path)
    else:
        print("No file selected.")